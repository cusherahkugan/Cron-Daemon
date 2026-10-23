from pptx import Presentation

from pptx_builder import build_pptx

PARA = ("Sentence one explains the idea. Sentence two adds detail about how it works. "
        "Sentence three gives an example. Sentence four wraps it up.")


def _section(i, n=1, empty=False):
    body = "".join(
        f"<h2>Heading {j}</h2><p>{PARA}</p><ul><li>point a</li><li>point b</li><li>point c</li></ul>"
        f"<pre>FROM ubuntu\nRUN apt-get update\nCMD echo hi</pre>"
        f"<table><tr><th>Name</th><th>Use</th></tr><tr><td>a</td><td>first</td></tr><tr><td>b</td><td>second</td></tr></table>"
        for j in range(n))
    return {"topic": f"T{i}", "subtopic": f"Demo {i}", "url": "https://example.com",
            "html": "" if empty else body}


def _words(prs):
    out = []
    for slide in prs.slides:
        out.append(sum(len(sh.text_frame.text.split()) for sh in slide.shapes if sh.has_text_frame))
    return out


def test_pptx_is_created(tmp_path):
    out = tmp_path / "t.pptx"
    n = build_pptx("2026-10-22", [_section(1, 2)], str(out))
    assert n == len(Presentation(str(out)).slides) >= 3   # title + agenda + content


def test_big_day_lands_in_target_range(tmp_path):
    out = tmp_path / "t.pptx"
    n = build_pptx("2026-10-22", [_section(i, 40) for i in range(15)], str(out), 50, 60)
    assert 50 <= n <= 60


def test_small_day_is_not_padded(tmp_path):
    out = tmp_path / "t.pptx"
    n = build_pptx("2026-10-22", [_section(i, 1) for i in range(2)], str(out), 50, 60)
    assert n < 50


def test_slides_are_not_stacked(tmp_path):
    out = tmp_path / "t.pptx"
    build_pptx("2026-10-22", [_section(i, 40) for i in range(15)], str(out))
    assert max(_words(Presentation(str(out)))) <= 230


def test_missing_article_gets_placeholder(tmp_path):
    out = tmp_path / "t.pptx"
    n = build_pptx("2026-10-22", [_section(1, empty=True), _section(2, 2)], str(out))
    texts = " ".join(sh.text_frame.text for s in Presentation(str(out)).slides
                     for sh in s.shapes if sh.has_text_frame)
    assert "No content could be fetched today." in texts and n >= 4