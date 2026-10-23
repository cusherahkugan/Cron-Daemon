"""Lecture slide deck (.pptx) built from the same notes as the PDF.

A day's notes are far too long to fit 50-60 slides word for word, so the deck is a
*condensed* version of the PDF: every topic keeps its headings, key sentences, bullet
points, a few code samples, tables and diagrams, trimmed just enough to fit the topic's
share of the slide budget. Slides are kept deliberately airy (only ~70% of the body area
is ever filled). The full text stays in the PDF.
"""
import base64
import copy
import io
import math
import re
import textwrap

from lxml import etree
from lxml import html as lh
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from condense import full_content
from scrape import embed_image

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN_X = 0.8
BODY_TOP = 1.6
BODY_BOTTOM = 6.85
BODY_W = SLIDE_W - 2 * MARGIN_X
BODY_H = BODY_BOTTOM - BODY_TOP

# (body font pt, fraction of the body area allowed to fill). Most spacious last; the
# first layout that reaches min_slides is used.
LAYOUTS = [(20, 0.72), (22, 0.66), (24, 0.60)]

GREEN = RGBColor(0x1A, 0x5F, 0x2A)
BLUE = RGBColor(0x0B, 0x3D, 0x91)
INK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x77, 0x77, 0x77)
RULE = RGBColor(0xDD, 0xDD, 0xDD)
CODE_BG = RGBColor(0xF4, 0xF4, 0xF4)
ROW_ALT = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# How much each topic is trimmed, from none (0) to headings + one key sentence (5).
#   p_sent/p_chars: sentences / characters kept per paragraph   p_group: paragraphs kept per heading
#   li_max/li_chars: bullets kept per list / characters per bullet
#   code_*, imgs, tables, table_rows: per-topic quotas (first N kept), code lines per block, rows per table
_LEVELS = [
    dict(p_sent=None, p_chars=None, li_max=None, li_chars=None, code_lines=None, code_blocks=None, imgs=None, tables=None, table_rows=None, p_group=None),
    dict(p_sent=3, p_chars=480, li_max=10, li_chars=220, code_lines=14, code_blocks=None, imgs=4, tables=4, table_rows=8, p_group=None),
    dict(p_sent=2, p_chars=320, li_max=7, li_chars=170, code_lines=10, code_blocks=6, imgs=3, tables=2, table_rows=6, p_group=None),
    dict(p_sent=1, p_chars=220, li_max=5, li_chars=140, code_lines=8, code_blocks=4, imgs=2, tables=1, table_rows=5, p_group=3),
    dict(p_sent=1, p_chars=200, li_max=3, li_chars=120, code_lines=6, code_blocks=2, imgs=1, tables=0, table_rows=0, p_group=1),
    dict(p_sent=1, p_chars=160, li_max=2, li_chars=100, code_lines=5, code_blocks=1, imgs=0, tables=0, table_rows=0, p_group=1),
]

_SENT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(])")
_DATA = re.compile(r"data:image/\w+;base64,(.*)", re.S)
CODE_REL = 0.78      # a code line costs this much of a body line (smaller font)
CODE_PAD = 0.8       # padding of a code box, in body lines
SOLO = 0.7           # images / tables take this share of a slide


# ------------------------------------------------------------------ layout maths
class _Lay:
    def __init__(self, fs: float, fill: float):
        self.fs = fs
        self.cap = BODY_H * fill / (fs * 1.2 / 72)          # body lines per slide
        self.cpl = int(BODY_W * 72 / (fs * 0.47))            # characters per line (Calibri)
        self.cpl_li = int((BODY_W - 0.35) * 72 / (fs * 0.47))
        self.cpl_h = int(BODY_W * 72 / (fs * 1.1 * 0.52))    # bold headings
        self.cfs = fs * 0.75                                 # code font
        self.ccpl = int((BODY_W - 0.3) * 72 / (self.cfs * 0.6))


def _lines(text: str, cpl: int) -> int:
    """Rendered line count of wrapped prose."""
    return max(1, len(textwrap.wrap(text, cpl)))


def _clines(line: str, cpl: int) -> int:
    """Rendered line count of one code line (indentation kept, hard wrap)."""
    return max(1, math.ceil(len(line) / cpl))


def _clip(s: str, m: int | None) -> str:
    if not m or len(s) <= m:
        return s
    return s[:m].rsplit(" ", 1)[0].rstrip(",;:-") + "…"


def _shorten(text: str, n_sent: int | None, max_chars: int | None) -> str:
    sents = _SENT.split(text)
    if n_sent:
        sents = sents[:n_sent]
    return _clip(" ".join(sents), max_chars)


# ------------------------------------------------------------------------ parsing
def _txt(el) -> str:
    return " ".join("".join(el.itertext()).split())


def _parse(body_html: str) -> list[tuple]:
    """Cleaned article HTML -> flat list of ("h"|"p"|"li"|"code"|"img"|"table", ...)."""
    if not body_html.strip():
        return []
    try:
        root = lh.fragment_fromstring(body_html, create_parent="div")
    except Exception:
        return []
    items: list[tuple] = []

    def add_img(el):
        m = _DATA.match(el.get("src") or "")
        if not m:
            return
        try:
            data = base64.b64decode(m.group(1))
            w, h = Image.open(io.BytesIO(data)).size
        except Exception:
            return
        items.append(("img", data, w, h))

    def add_list(el):
        for li in el.xpath("./li"):
            own = copy.deepcopy(li)
            for sub in own.xpath("./ul|./ol"):
                sub.drop_tree()
            t = _txt(own)
            if t:
                items.append(("li", t))
            for im in own.iter("img"):
                add_img(im)
            for sub in li.xpath("./ul|./ol"):
                add_list(sub)

    def walk(el):
        for c in el:
            if not isinstance(c.tag, str):
                continue
            t = c.tag
            if t in ("h4", "h5"):
                s = _txt(c)
                if s:
                    items.append(("h", s))
            elif t == "p":
                s = _txt(c)
                if s:
                    items.append(("p", s))
                for im in c.iter("img"):
                    add_img(im)
            elif t in ("ul", "ol"):
                add_list(c)
            elif t == "pre":
                text = "".join(c.itertext()).strip("\n")
                if text.strip():
                    items.append(("code", text))
            elif t == "table":
                rows = []
                for tr in c.xpath(".//tr"):
                    cells = [_txt(td) for td in tr.xpath("./th|./td")]
                    if any(cells):
                        rows.append(cells)
                if len(rows) >= 2:
                    items.append(("table", rows))
            elif t == "img":
                add_img(c)
            else:
                walk(c)

    walk(root)
    return items


# ------------------------------------------------------------------- condensing
def _condense(items: list[tuple], lv: int) -> list[tuple]:
    P = _LEVELS[lv]
    out: list[tuple] = []
    n_p = n_li = n_code = n_img = n_tab = 0
    for it in items:
        k = it[0]
        if k != "li":
            n_li = 0
        if k == "h":
            out.append(it)
            n_p = 0
        elif k == "p":
            if P["p_group"] is not None and n_p >= P["p_group"]:
                continue
            n_p += 1
            out.append(("p", _shorten(it[1], P["p_sent"], P["p_chars"])))
        elif k == "li":
            n_li += 1
            if P["li_max"] is not None and n_li > P["li_max"]:
                continue
            out.append(("li", _clip(it[1], P["li_chars"])))
        elif k == "code":
            if P["code_blocks"] is not None and n_code >= P["code_blocks"]:
                continue
            n_code += 1
            lines = it[1].split("\n")
            if P["code_lines"] is not None and len(lines) > P["code_lines"]:
                lines = lines[:P["code_lines"]] + ["…"]
            out.append(("code", "\n".join(lines)))
        elif k == "img":
            if P["imgs"] is not None and n_img >= P["imgs"]:
                continue
            n_img += 1
            out.append(it)
        elif k == "table":
            if P["tables"] is not None and n_tab >= P["tables"]:
                continue
            n_tab += 1
            rows = it[1] if P["table_rows"] is None else it[1][:P["table_rows"] + 1]
            out.append(("table", rows))
    # a heading with nothing under it is just noise
    return [x for i, x in enumerate(out)
            if not (x[0] == "h" and (i + 1 == len(out) or out[i + 1][0] == "h"))]


def _split_text(text: str, max_chars: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    pieces: list[str] = []
    cur = ""
    for s in _SENT.split(text):
        while len(s) > max_chars:
            cut = s.rfind(" ", 0, max_chars)
            cut = cut if cut > 0 else max_chars
            head, s = s[:cut], s[cut:].lstrip()
            if cur:
                pieces.append(cur)
                cur = ""
            pieces.append(head)
        if cur and len(cur) + 1 + len(s) > max_chars:
            pieces.append(cur)
            cur = s
        else:
            cur = f"{cur} {s}".strip()
    if cur:
        pieces.append(cur)
    return pieces


def _weights(rows: list[list[str]]) -> list[int]:
    """Relative column widths of a table, from its longest cell per column."""
    return [min(40, max(len(r[c]) for r in rows)) + 6 for c in range(len(rows[0]))]


def _row_lines(row: list[str], weights: list[int], lay: _Lay) -> int:
    """Lines needed by the tallest cell of a table row (table text is 0.8x body size)."""
    return max(_lines(c, max(6, int(lay.cpl / 0.8 * 0.92 * weights[i] / sum(weights))))
               for i, c in enumerate(row))


def _units(items: list[tuple], lay: _Lay) -> list[dict]:
    """Items -> sized units, each small enough to sit on one slide."""
    cap = lay.cap
    out: list[dict] = []
    for it in items:
        k = it[0]
        if k == "h":
            out.append({"kind": "h", "text": it[1], "cost": _lines(it[1], lay.cpl_h) * 1.1 + 0.9})
        elif k in ("p", "li"):
            cpl = lay.cpl_li if k == "li" else lay.cpl
            for piece in _split_text(it[1], int(cap * 0.6 * cpl)):
                out.append({"kind": k, "text": piece, "cost": _lines(piece, cpl) + 0.4})
        elif k == "code":
            lines = [ln.expandtabs(4).rstrip() for ln in it[1].split("\n")]
            room = (cap * 0.85 - CODE_PAD) / CODE_REL
            chunk: list[str] = []
            used = 0.0
            for ln in lines + [None]:
                c = _clines(ln, lay.ccpl) if ln is not None else 0
                if chunk and (ln is None or used + c > room):
                    out.append({"kind": "code", "lines": chunk, "cost": used * CODE_REL + CODE_PAD})
                    chunk, used = [], 0.0
                if ln is not None:
                    chunk.append(ln)
                    used += c
        elif k == "img":
            out.append({"kind": "img", "data": it[1], "w": it[2], "h": it[3], "cost": cap * SOLO, "solo": True})
        elif k == "table":
            rows = it[1]
            ncols = max(len(r) for r in rows)
            rows = [r + [""] * (ncols - len(r)) for r in rows]
            weights = _weights(rows)

            def row_cost(r):
                return _row_lines(r, weights, lay) * 0.8 + 0.35

            hdr, body = rows[0], rows[1:]
            chunk, used = [], row_cost(hdr)
            for r in body + [None]:
                c = row_cost(r) if r is not None else 0
                if chunk and (r is None or used + c > cap * SOLO):
                    out.append({"kind": "table", "rows": [hdr] + chunk, "cost": cap * SOLO, "solo": True})
                    chunk, used = [], row_cost(hdr)
                if r is not None:
                    chunk.append(r)
                    used += c
    return out


def _paginate(units: list[dict], lay: _Lay) -> list[list[dict]]:
    slides: list[list[dict]] = []
    cur: list[dict] = []
    used = 0.0

    def flush():
        nonlocal cur, used
        if cur:
            slides.append(cur)
        cur, used = [], 0.0

    for i, u in enumerate(units):
        c = u["cost"]
        if u["kind"] == "h":
            # keep a heading together with the block that follows it
            nxt = units[i + 1]["cost"] if i + 1 < len(units) and units[i + 1]["kind"] != "h" else 0
            if cur and used + c + nxt > lay.cap:
                flush()
        elif u.get("solo"):
            if used + c > lay.cap:
                flush()
        elif cur and used + c > lay.cap:
            flush()
        cur.append(u)
        used += c
        if u.get("solo"):
            flush()
    flush()
    if len(slides) > 1 and all(x["kind"] == "h" for x in slides[-1]):
        slides.pop()
    return slides


def _slides(items: list[tuple], lv: int, lay: _Lay) -> list[list[dict]]:
    slides = _paginate(_units(_condense(items, lv), lay), lay)
    if not slides:
        text = "See the PDF for the full notes."
        slides = [[{"kind": "p", "text": text, "cost": 1.4}]]
    return slides


def _thin(items: list[tuple], lay: _Lay, budget: int) -> list[list[dict]]:
    """Last resort for very long articles: keep the most condensed form of evenly spaced
    heading groups (so the whole article is still represented, not just its start)."""
    last = len(_LEVELS) - 1
    groups: list[list[tuple]] = [[]]
    for it in items:
        if it[0] == "h":
            groups.append([])
        groups[-1].append(it)
    lead, rest = groups[0], groups[1:]
    slides = _slides(items, last, lay)
    for keep in range(len(rest), 0, -1):
        picked = [rest[min(len(rest) - 1, int(i * len(rest) / keep))] for i in range(keep)]
        subset = lead + [it for g in picked for it in g]
        slides = _slides(subset, last, lay)
        if len(slides) <= budget:
            break
    return slides[:budget]


def _allocate(needs: list[int], avail: int) -> list[int]:
    """Share `avail` slides between topics: small topics keep what they need, the rest
    split what is left evenly."""
    alloc: list = [None] * len(needs)
    todo = list(range(len(needs)))
    remaining = avail
    while todo:
        share = remaining // len(todo)
        small = [i for i in todo if needs[i] <= share]
        if not small:
            break
        for i in small:
            alloc[i] = needs[i]
            remaining -= needs[i]
            todo.remove(i)
    if todo:
        share, extra = divmod(remaining, len(todo))
        for k, i in enumerate(todo):
            alloc[i] = max(1, share + (1 if k < extra else 0))
    return alloc


def _plan(sections: list[dict], raw: list[list[tuple]], lay: _Lay, avail: int) -> list[dict]:
    needs = [len(_slides(items, 0, lay)) for items in raw]
    alloc = _allocate(needs, avail)
    out = []
    for s, items, budget in zip(sections, raw, alloc):
        slides = []
        for lv in range(len(_LEVELS)):
            slides = _slides(items, lv, lay)
            if len(slides) <= budget:
                break
        if len(slides) > budget:
            slides = _thin(items, lay, budget)
        for i, units in enumerate(slides, 1):
            out.append({"type": "content", "topic": s["topic"], "subtopic": s["subtopic"],
                        "url": s.get("url", ""), "part": i, "parts": len(slides), "units": units})
    return out


# ---------------------------------------------------------------------- rendering
def _tb(slide, x, y, w, h):
    tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _run(p, text, size, color=INK, bold=False, font="Calibri"):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = font
    r.font.color.rgb = color


def _rect(slide, x, y, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _bullet(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.35))))
    pPr.set("indent", str(-int(Inches(0.3))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", "1A5F2A")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def _footer(slide, url, n, total):
    if url:
        _run(_tb(slide, MARGIN_X, 7.02, 9.5, 0.3).paragraphs[0], _clip(url, 110), 9, GREY)
    p = _tb(slide, SLIDE_W - MARGIN_X - 1.5, 7.02, 1.5, 0.3).paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, f"{n} / {total}", 10, GREY)


def _text_group(slide, group, y, lay):
    h = sum(u["cost"] for u in group) * lay.fs * 1.2 / 72
    tf = _tb(slide, MARGIN_X, y, BODY_W, h)
    for i, u in enumerate(group):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if u["kind"] == "h":
            p.space_before = Pt(lay.fs * 0.5) if i else Pt(0)
            p.space_after = Pt(lay.fs * 0.2)
            _run(p, u["text"], lay.fs * 1.1, BLUE, bold=True)
        else:
            p.space_after = Pt(lay.fs * 0.35)
            _run(p, u["text"], lay.fs)
            if u["kind"] == "li":
                _bullet(p)
    return y + h + 0.1


def _code_box(slide, u, y, lay):
    n = sum(_clines(ln, lay.ccpl) for ln in u["lines"])
    h = n * lay.cfs * 1.2 / 72 + 0.3
    box = _rect(slide, MARGIN_X, y, BODY_W, h, CODE_BG)
    _rect(slide, MARGIN_X, y, 0.06, h, GREEN)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = (Inches(0.2), Inches(0.12), Inches(0.12), Inches(0.1))
    for i, ln in enumerate(u["lines"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        _run(p, ln or " ", lay.cfs, INK, font="Consolas")
    return y + h + 0.12


def _picture(slide, u, y):
    avail = BODY_BOTTOM - y - 0.05
    ratio = u["w"] / u["h"]
    w_in = min(BODY_W, avail * ratio, u["w"] / 80)
    h_in = w_in / ratio
    slide.shapes.add_picture(io.BytesIO(u["data"]), Inches(MARGIN_X + (BODY_W - w_in) / 2), Inches(y),
                             Inches(w_in), Inches(h_in))
    return y + h_in + 0.1


def _table(slide, u, y, lay):
    rows = u["rows"]
    nr, nc = len(rows), max(len(r) for r in rows)
    rows = [r + [""] * (nc - len(r)) for r in rows]
    weights = _weights(rows)
    fs = lay.fs * 0.8
    heights = [max(0.4, _row_lines(r, weights, lay) * fs * 1.2 / 72 + 0.16) for r in rows]
    shape = slide.shapes.add_table(nr, nc, Inches(MARGIN_X), Inches(y), Inches(BODY_W), Inches(sum(heights)))
    tbl = shape.table
    for c in range(nc):
        tbl.columns[c].width = Inches(BODY_W * weights[c] / sum(weights))
    for r in range(nr):
        tbl.rows[r].height = Inches(heights[r])
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = GREEN if r == 0 else (ROW_ALT if r % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            _run(cell.text_frame.paragraphs[0], rows[r][c], fs, WHITE if r == 0 else INK, bold=(r == 0))
    return y + sum(heights) + 0.1


def _content_slide(prs, spec, n, total, lay):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_W, 0.14, GREEN)
    tag = spec["topic"].upper() + (f"  ·  {spec['part']}/{spec['parts']}" if spec["parts"] > 1 else "")
    _run(_tb(slide, MARGIN_X, 0.38, BODY_W, 0.3).paragraphs[0], tag, 12, GREEN, bold=True)
    _run(_tb(slide, MARGIN_X, 0.66, BODY_W, 0.6).paragraphs[0], _clip(spec["subtopic"], 70), 28, INK, bold=True)
    _rect(slide, MARGIN_X, 1.36, BODY_W, 0.02, RULE)

    y = BODY_TOP
    group: list[dict] = []
    for u in spec["units"]:
        if u["kind"] in ("h", "p", "li"):
            group.append(u)
            continue
        if group:
            y = _text_group(slide, group, y, lay)
            group = []
        if u["kind"] == "code":
            y = _code_box(slide, u, y, lay)
        elif u["kind"] == "img":
            y = _picture(slide, u, y)
        elif u["kind"] == "table":
            y = _table(slide, u, y, lay)
    if group:
        _text_group(slide, group, y, lay)
    _footer(slide, spec["url"], n, total)


def _title_slide(prs, date_str, n_topics):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, GREEN)
    _run(_tb(slide, 1.0, 2.4, 11.3, 1.2).paragraphs[0], "Daily Study Notes", 54, WHITE, bold=True)
    _run(_tb(slide, 1.0, 3.6, 11.3, 0.6).paragraphs[0], date_str, 28, RGBColor(0xCF, 0xE8, 0xD4))
    _run(_tb(slide, 1.0, 4.4, 11.3, 0.5).paragraphs[0],
         f"Lecture slides  ·  {n_topics} topics  ·  condensed from the full PDF notes", 18, RGBColor(0xCF, 0xE8, 0xD4))


def _agenda_slide(prs, sections, n, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, SLIDE_W, 0.14, GREEN)
    _run(_tb(slide, MARGIN_X, 0.55, BODY_W, 0.7).paragraphs[0], "Today's topics", 34, INK, bold=True)
    _rect(slide, MARGIN_X, 1.36, BODY_W, 0.02, RULE)
    cols = 2 if len(sections) <= 16 else 3
    per = math.ceil(len(sections) / cols)
    fs = 16 if cols == 2 else 14
    col_w = BODY_W / cols - 0.2
    for c in range(cols):
        chunk = sections[c * per:(c + 1) * per]
        if not chunk:
            continue
        tf = _tb(slide, MARGIN_X + c * (BODY_W / cols), BODY_TOP + 0.1, col_w, 4.8)
        for i, s in enumerate(chunk):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(fs * 0.6)
            _run(p, s["topic"], fs, GREEN, bold=True)
            _run(p, "  " + _clip(s["subtopic"], int(col_w * 72 / (fs * 0.52)) - len(s["topic"]) - 4), fs)
    _footer(slide, "", n, total)


def build_pptx(date_str: str, sections: list[dict], out_path: str,
               min_slides: int = 50, max_slides: int = 60) -> int:
    """Write the lecture deck and return its slide count. Uses the same `sections`
    (topic / subtopic / url / html) as build_pdf."""
    raw = []
    for s in sections:
        body = full_content(s.get("html", ""),
                            embed=lambda src, base=s.get("url", ""): embed_image(src, base))
        raw.append(_parse(body) or [("p", "No content could be fetched today.")])

    lay, content = None, []
    for fs, fill in LAYOUTS:
        lay = _Lay(fs, fill)
        content = _plan(sections, raw, lay, max_slides - 2)
        if len(content) + 2 >= min_slides:
            break
    total = len(content) + 2
    if total < min_slides:
        print(f"NOTE: only {total} slides - today's notes are shorter than the {min_slides}-slide target")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    _title_slide(prs, date_str, len(sections))
    _agenda_slide(prs, sections, 2, total)
    for i, spec in enumerate(content, 3):
        _content_slide(prs, spec, i, total, lay)
    prs.save(out_path)
    return len(prs.slides)